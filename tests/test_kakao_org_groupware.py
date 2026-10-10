# -*- coding: utf-8 -*-
"""사내 그룹웨어와 같은 조직도 구성 (그룹사 › 한국지주 › 케이씨에스 › 직속 구성원 + 부서 10개, 관리자 숨김)"""

import os
import sys
import unittest
from unittest import mock

from tests.kakao_fakes import make_demo_state, app

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'scripts'))
import import_org_xlsx as imp  # noqa: E402

from client.kakao.org_tab import OrgTab, MembersDialog, CompactMemberRow, CompanyHeader  # noqa: E402


def rows(lst):
    return [lst.item(i).data(0x0100) for i in range(lst.count())]


class GroupwarePathTest(unittest.TestCase):
    def test_flatten_to_departments(self):
        f = imp.to_groupware_path
        self.assertEqual(f(['임원실']), [])
        self.assertEqual(f(['관리팀']), ['관리팀'])
        self.assertEqual(f(['MICE 사업팀']), ['MICE사업팀'])
        self.assertEqual(f(['영업부', '영업1본부', '1팀(보안)']), ['영업1부', '1팀(보안)'])
        self.assertEqual(f(['기술연구소', '연구2실', '3팀(DB)']), ['기술연구소2실', '3팀(DB)'])

    def test_layout_keeps_every_concurrent_team_and_head_seat(self):
        """(*) 겸직은 같은 부서의 다른 팀도 각각 남고, 팀 안의 부서장 직함은 부서 바로 아래에도 둔다"""
        users = [{'username': '테스트임원', 'company': '케이씨에스', 'department': '1팀(Oracle)', 'title': '상무이사',
                  'org_path': ['영업부', '영업3본부', '1팀(Oracle)'], 'account': 'smyoon', 'extension': '5142',
                  'mobile': '', 'concurrent': [{'org_path': ['기술연구소', '연구2실', '1팀(Oracle)'], 'title': '상무이사'},
                                               {'org_path': ['기술연구소', '연구2실', '3팀(DB)'], 'title': '상무이사'}]}]
        out = imp.apply_groupware_layout(users)
        ys = out[0]
        self.assertEqual(ys['org_path'], ['영업3부'])                         # 상무이사 → 영업3부 부서장 자리
        self.assertEqual([c['org_path'] for c in ys['concurrent']],
                         [['영업3부', '1팀(Oracle)'], ['기술연구소2실'], ['기술연구소2실', '1팀(Oracle)'],
                          ['기술연구소2실', '3팀(DB)']])
        self.assertTrue(all(c['title'] == '상무이사' for c in ys['concurrent']))
        self.assertEqual(out[-1]['username'], 'KCS관리자')
        self.assertEqual(out[-1]['account'], 'kcs_administrator')
        self.assertEqual(out[-1]['org_path'], [])

    def test_unit_rename_keeps_team_level(self):
        """엑셀의 '2팀(SI)' 은 'AI/AX 솔루션사업팀' 으로 바뀌고, 이름이 숫자로 시작하지 않아도 연구1실 아래 팀으로 남는다"""
        import openpyxl
        ws = openpyxl.Workbook().active
        for r, row in enumerate([['기술연구소(t-research)'], ['연구1실 (t-research1)'], ['2팀(SI)'],
                                 ['수석', '도승현', 'doshyun', '5101', '010-0000-0000']], start=1):
            for c, value in enumerate(row, start=1):
                ws.cell(r, c, value)
        entries, units = imp.parse_sheet(ws)
        self.assertEqual(entries[0]['path'], ['기술연구소', '연구1실', 'AI/AX 솔루션사업팀'])
        self.assertEqual(imp.to_groupware_path(entries[0]['path']), ['기술연구소1실', 'AI/AX 솔루션사업팀'])
        self.assertNotIn(['기술연구소', '연구1실', '2팀(SI)'], units)

    def test_department_order_matches_groupware(self):
        self.assertEqual(imp.GROUPWARE_DEPTS, ['관리팀', '마케팅팀', '신사업팀', 'MICE사업팀', '영업1부', '영업2부',
                                               '영업3부', '기술연구소1실', '기술연구소2실', '기술연구소3실'])


# 부서별 인원(겸직 포함, 사람 기준). 겸직자도 그 부서 인원에 센다 - 엑셀 (*) 와 팀 안 부서장 규칙 반영
# 그룹웨어 화면은 관리팀 5 · MICE사업팀 2 지만, 화면에 이름이 없어 MICE사업팀의 누가 옮겼는지 알 수 없으므로
# 엑셀 소속 그대로 둔다 (류병두: 엑셀 두 버전 모두 MICE 사업팀) → 관리팀 4 · MICE사업팀 3
IMAGE_COUNTS = {'관리팀': 4, '마케팅팀': 3, '신사업팀': 3, 'MICE사업팀': 3, '영업1부': 2, '영업2부': 4, '영업3부': 5,
                '기술연구소1실': 9, '기술연구소2실': 14, '기술연구소3실': 10}
XLSX = os.path.expanduser('~/Downloads/연락처 및 조직도(20260901).xlsx')


class GroupwareOverridesTest(unittest.TestCase):
    def test_overrides_move_and_drop(self):
        users = [
            {'username': '류병두', 'company': 'K', 'department': 'MICE 사업팀', 'title': '매니져',
             'org_path': ['MICE 사업팀'], 'account': 'rbd', 'extension': '', 'mobile': '', 'concurrent': []},
            {'username': '보정대상', 'company': 'K', 'department': 'MICE 사업팀', 'title': '매니져',
             'org_path': ['MICE 사업팀'], 'account': 'x', 'extension': '', 'mobile': '', 'concurrent': []},
            {'username': '엄태식', 'company': 'K', 'department': '1팀(Tandem)', 'title': '실장',
             'org_path': ['기술연구소', '연구1실', '1팀(Tandem)'], 'account': 'tseom', 'extension': '', 'mobile': '',
             'concurrent': [{'org_path': ['신사업팀'], 'title': '실장'},
                            {'org_path': ['기술연구소', '연구1실', '2팀(SI)'], 'title': '실장'}]},
        ]
        with mock.patch.dict(imp.GROUPWARE_OVERRIDES, {'보정대상': {'dept': '관리팀'}}):
            out = {u['username']: u for u in imp.apply_groupware_layout(users)}
        self.assertNotIn('류병두', imp.GROUPWARE_OVERRIDES)                   # 근거 없던 보정은 삭제 (엑셀 소속 그대로)
        self.assertEqual(out['류병두']['org_path'], ['MICE사업팀'])
        self.assertEqual(out['보정대상']['org_path'], ['관리팀'])               # 확인된 사람은 보정으로 옮길 수 있다
        self.assertEqual(out['엄태식']['org_path'], ['신사업팀'])
        # 주소속은 신사업팀으로 옮겨도 연구1실 실장 자리와 1팀·2팀 겸직은 남는다
        self.assertEqual([c['org_path'] for c in out['엄태식']['concurrent']],
                         [['기술연구소1실'], ['기술연구소1실', '1팀(Tandem)'], ['기술연구소1실', '2팀(SI)']])
        self.assertEqual(out['보정대상']['concurrent'], [])                     # 소속 이동은 겸직이 아님

    @unittest.skipUnless(os.path.isfile(XLSX), '원본 엑셀 없음')
    def test_real_excel_matches_groupware_image(self):
        """원본 엑셀을 변환한 부서별 숫자 (그룹웨어 화면과는 관리팀·MICE사업팀만 한 명씩 다름 - 위 IMAGE_COUNTS 설명)"""
        import warnings
        warnings.filterwarnings('ignore')
        import openpyxl
        ws = openpyxl.load_workbook(XLSX, data_only=True).worksheets[0]
        entries, _units = imp.parse_sheet(ws)
        users = imp.apply_groupware_layout(imp.build_users(entries, '케이씨에스'))
        people, in_depts = {}, set()
        for u in users:
            for path in [u['org_path']] + [c['org_path'] for c in u['concurrent']]:
                if path:
                    people.setdefault(path[0], set()).add(u['username'])
                    in_depts.add(u['username'])
        self.assertEqual({d: len(v) for d, v in people.items()}, IMAGE_COUNTS)
        by_name = {u['username']: u for u in users}
        places = lambda n: {tuple(by_name[n]['org_path'])} | {tuple(c['org_path']) for c in by_name[n]['concurrent']}  # noqa: E731
        # 엑셀의 (*) 겸직자 5명은 모든 조직에 나온다
        self.assertEqual(places('이동현'), {('마케팅팀',), ('신사업팀',)})
        self.assertEqual(places('문동민'), {('영업2부', '1팀(System)'), ('신사업팀',)})
        self.assertEqual(places('조종익'), {('기술연구소3실', '2팀(Chip)'), ('영업1부', '1팀(보안)')})
        self.assertTrue({('기술연구소1실', '1팀(Tandem)'), ('기술연구소1실', 'AI/AX 솔루션사업팀'),
                         ('신사업팀',)} <= places('엄태식'))
        teams = {tuple(p) for u in users for p in [u['org_path']] + [c['org_path'] for c in u['concurrent']]}
        self.assertNotIn(('기술연구소1실', '2팀(SI)'), teams)                       # 예전 이름은 남지 않는다
        self.assertTrue({('영업3부', '1팀(Oracle)'), ('기술연구소2실', '1팀(Oracle)'),
                         ('기술연구소2실', '3팀(DB)')} <= places('윤성목'))
        self.assertIn(('영업2부',), places('김성준'))                               # 본부장 → 부서장 자리
        self.assertEqual(places('류병두'), {('MICE사업팀',)})                         # 엑셀 소속 그대로
        self.assertEqual(len(in_depts), 52)
        roots = [u['username'] for u in users if not u['org_path']]
        self.assertEqual(sorted(roots), sorted(['김광묵', 'KCS관리자']))


class GroupwareOrgTabTest(unittest.TestCase):
    def setUp(self):
        app()
        self.state, self.conn = make_demo_state()
        self.state.org = {'holding': {'name': '한국지주'},
                          'companies': [{'name': '한네트'},
                                        {'name': '케이씨에스', 'group_name': '케이씨에스',
                                         'depth': 1, 'root_members': ['홍길동', 'KCS관리자'],
                                         'hidden_members': ['KCS관리자'],
                                         'count_excludes_root': True,
                                         'units': [[d] for d in imp.GROUPWARE_DEPTS]}]}
        people = (('홍길동', [], '대표이사', []), ('김철수', ['관리팀'], '팀장', []),
                  ('이영희', ['기술연구소1실', '2팀(SI)'], '수석', [{'org_path': ['신사업팀'], 'title': '수석'}]),
                  ('박민수', ['영업1부', '1팀(보안)'], '부장', []))
        for name, path, title, extra in people:
            self.conn.push('user_updated', user=dict(self.state.user(name), company='케이씨에스', org_path=path,
                                                    department=path[0] if path else '', title=title,
                                                    concurrent=extra, extension='5100', account=name))
        self.tab = OrgTab(self.state, self.state.settings)

    def add_admin(self):
        self.conn.push('user_updated', user=dict(self.state.user('최지우'), username='KCS관리자',
                                                 display_name='KCS관리자', company='케이씨에스', org_path=[],
                                                 department='', title='', concurrent=[], account='kcs_administrator'))
        self.tab.refresh()

    def test_tree_shape_like_groupware(self):
        data = rows(self.tab.list)
        # 케이씨에스는 그룹사 아래 한국지주 · 한네트와 같은 단계 자기 순서 자리에, 법인명 단계 없이
        self.assertEqual(data[:3], [('section', '그룹사'), ('company', '한국지주'), ('company', '한네트')])
        idx = data.index(('company', '케이씨에스'))
        self.assertGreater(idx, data.index(('company', '한네트')))
        self.assertNotIn(('group', '케이씨에스'), data)
        self.assertEqual(data[idx + 1], ('member', '홍길동'))                 # 회사 직속(대표이사) 먼저
        depts = [d[2][0] for d in data if d[0] == 'unit']
        self.assertEqual(depts, imp.GROUPWARE_DEPTS)                         # 10개 부서, 그룹웨어 순서
        self.assertNotIn(('member', '김철수'), data)                          # 부서는 처음에 접혀 있음
        row = self.tab.list.itemWidget(self.tab.list.find_data(('company', '케이씨에스')))
        self.assertEqual(row.title.text(), '케이씨에스')
        self.assertIsInstance(row, CompanyHeader)                            # 다른 계열사처럼 건물 아이콘 머리글
        self.assertIsNotNone(row.chip)                                       # + 인원 칩
        titles = [getattr(self.tab.list.itemWidget(self.tab.list.item(i)), 'title', None)
                  for i in range(self.tab.list.count())]
        self.assertNotIn('주식회사케이씨에스', [t.text() for t in titles if t is not None])   # 법인명 단계 없음

    def test_root_members_and_dept_arrows_align_with_company_title(self):
        """대표이사 · 관리자 사진과 부서 › 화살표 시작이 '케이씨에스' 글자 시작 줄에 맞는다 (2026-10-10 요청)"""
        from client.kakao.org_tab import TreeNodeRow
        self.add_admin()
        self.state.me = 'KCS관리자'
        self.tab.refresh()
        lst = self.tab.list
        widget = lambda key: lst.itemWidget(lst.find_data(key))
        for w in (widget(('company', '케이씨에스')), widget(('member', '홍길동')), widget(('member', 'KCS관리자'))):
            w.layout().activate()
        title_x = widget(('company', '케이씨에스')).title.geometry().left()
        for name in ('홍길동', 'KCS관리자'):
            self.assertEqual(widget(('member', name)).avatar.geometry().left(), title_x)
        dept = widget(('unit', '케이씨에스', ('관리팀',)))
        dept.layout().activate()
        chevron = dept.layout().itemAt(0).widget()
        self.assertEqual(chevron.geometry().left() + TreeNodeRow.INK, title_x)

    def test_admin_hidden_from_others(self):
        self.add_admin()
        data = rows(self.tab.list)
        self.assertNotIn(('member', 'KCS관리자'), data)
        self.tab.search_edit.setText('kcs_administrator')                     # 검색해도 안 나온다
        self.tab.refresh()
        self.assertNotIn(('member', 'KCS관리자'), rows(self.tab.list))

    def test_admin_login_sees_everything(self):
        self.add_admin()
        self.state.me = 'KCS관리자'
        self.tab.refresh()
        data = rows(self.tab.list)
        self.assertIn(('member', 'KCS관리자'), data)
        self.assertIn(('member', '홍길동'), data)

    def test_expand_department_and_concurrent(self):
        item = self.tab.list.find_data(('unit', '케이씨에스', ('관리팀',)))
        self.tab.list.itemClicked.emit(item)
        self.assertIn(('member', '김철수'), rows(self.tab.list))
        row = self.tab.list.itemWidget(self.tab.list.find_data(('member', '김철수')))
        self.assertIsInstance(row, CompactMemberRow)
        self.assertEqual(row.name.text(), '김철수 팀장')
        self.tab.list.itemClicked.emit(self.tab.list.find_data(('unit', '케이씨에스', ('신사업팀',))))
        data = rows(self.tab.list)
        idx = data.index(('unit', '케이씨에스', ('신사업팀',)))
        self.assertEqual(data[idx + 1], ('member', '이영희'))                  # 겸직자도 그 부서에 표시

    def test_department_members_start_at_department_name(self):
        from PyQt5.QtCore import QPoint
        self.tab.resize(400, 700)
        self.tab.show()
        key = ('unit', '케이씨에스', ('관리팀',))
        self.tab.list.itemClicked.emit(self.tab.list.find_data(key))
        data = rows(self.tab.list)
        dept = self.tab.list.itemWidget(self.tab.list.find_data(key))
        member = self.tab.list.itemWidget(self.tab.list.item(data.index(key) + 1))
        self.assertIsInstance(member, CompactMemberRow)
        for w in (dept, member):
            w.layout().activate()
        x = lambda w: w.mapTo(self.tab.list.viewport(), QPoint(0, 0)).x()   # noqa: E731
        self.assertEqual(x(member.avatar), x(dept.title))                     # 아바타 = 부서 이름 시작 위치
        self.tab.hide()

    def test_department_shows_head_then_teams_with_concurrent(self):
        """부서장(팀 안 겸직) → 팀 → 팀원. 같은 사람이 부서장 자리와 두 팀에 모두 '겸직'으로 나온다"""
        for name, path, title, extra in (
                ('최지우', ['기술연구소2실'], '상무이사', [{'org_path': ['기술연구소2실', '1팀(Oracle)'], 'title': '상무이사'},
                                                      {'org_path': ['기술연구소2실', '3팀(DB)'], 'title': '상무이사'}]),
                ('정하늘', ['기술연구소2실', '1팀(Oracle)'], '팀장', []),
                ('강민지', ['기술연구소2실', '3팀(DB)'], '팀장', [])):
            self.conn.push('user_updated', user=dict(self.state.user(name), company='케이씨에스', org_path=path,
                                                    department=path[0], title=title, concurrent=extra, account=name))
        self.tab.refresh()
        key = ('unit', '케이씨에스', ('기술연구소2실',))
        dept = self.tab.list.itemWidget(self.tab.list.find_data(key))
        self.assertEqual(dept.chip.count, 3)                                  # 사람 기준 (겸직은 한 번만)
        company = self.tab.list.itemWidget(self.tab.list.find_data(('company', '케이씨에스'))).chip.count
        # 총원: 부서 소속 6명(김철수·이영희·박민수·최지우·정하늘·강민지). 세 곳에 나오는 최지우도 1명
        self.assertEqual(company, 6)
        self.tab.list.itemClicked.emit(self.tab.list.find_data(key))
        data = rows(self.tab.list)
        i = data.index(key)
        self.assertEqual(data[i + 1:i + 7], [
            ('member', '최지우'),                                             # 부서장 자리
            ('unit', '케이씨에스', ('기술연구소2실', '1팀(Oracle)')), ('member', '최지우'), ('member', '정하늘'),
            ('unit', '케이씨에스', ('기술연구소2실', '3팀(DB)')), ('member', '최지우')])
        team_rows = [self.tab.list.itemWidget(self.tab.list.item(j)) for j in (i + 3, i + 6)]
        self.assertTrue(all(r.concurrent_title == '상무이사' for r in team_rows))      # 팀에서는 '겸직'
        team = ('unit', '케이씨에스', ('기술연구소2실', '1팀(Oracle)'))
        self.tab.list.itemClicked.emit(self.tab.list.find_data(team))           # 팀 접기
        self.assertNotIn(('member', '정하늘'), rows(self.tab.list))
        self.assertIn(key, rows(self.tab.list))                                 # 부서는 그대로 펼침

    def test_company_count_excludes_root_accounts(self):
        company = self.tab.list.itemWidget(self.tab.list.find_data(('company', '케이씨에스')))
        self.assertEqual(company.chip.count, 3)    # 김철수·이영희·박민수 (직속 홍길동 제외) - 그룹웨어 방식

    def test_count_chip_opens_member_list(self):
        node = self.tab.list.itemWidget(self.tab.list.find_data(('unit', '케이씨에스', ('기술연구소1실',))))
        self.assertEqual(node.chip.count, 1)
        dlg = self.tab.show_members('기술연구소1실', ['이영희'], exec_=False)
        self.assertIsInstance(dlg, MembersDialog)
        self.assertEqual(dlg.rows(), ['이영희'])
        dlg.search.setText('없는사람')
        self.assertEqual(dlg.rows(), [])

    def test_search_by_phone_or_id_expands(self):
        self.tab.search_edit.setText('박민수')
        data = rows(self.tab.list)
        self.assertIn(('member', '박민수'), data)
        self.assertEqual(self.tab.search_edit.placeholderText(), '이름, 직위, 직책, 직급, 부서, 전화, 아이디')


if __name__ == '__main__':
    unittest.main()
